"""Planilla Excel de la jornada."""

import pandas as pd
import pytest
from openpyxl import load_workbook

from proyecto_grado.dispatch.config import BloqueHorario
from proyecto_grado.piloto.planilla import (
    DatosPlanilla,
    escribir_planilla,
    horario_bloques,
    nombre_disponible,
)

BLOQUES = (BloqueHorario("madrugada", 4), BloqueHorario("pico_manana", 5))


def rec():
    return pd.DataFrame(
        {
            "route": [1, 1, 3],
            "dia_op": "2026-10-13",
            "hora": ["04:30", "05:00", "05:00"],
            "bloque": ["madrugada", "pico_manana", "pico_manana"],
            "regimen": "ordinario",
            "demanda_predicha": [12.34, 80.0, 60.0],
            "divisor": [20.0, 25.0, 30.0],
            "q3_ordinario": [30.0, 30.0, 30.0],
            "despachos_recomendados": [0, 3, 2],
            "despachos_regla_normal": [1, 3, 2],
            "en_disponibilidad": [False, True, True],
        }
    )


def datos(**k):
    base = {
        "fecha": "2026-10-13",
        "estado": "OK",
        "grupo": "SOMBRA",
        "id_ejecucion": "20261013_003000",
        "tipo_dia": "LABORAL",
        "regimen": "ordinario",
        "factores": {1: 0.87, 3: 0.91},
        "ultimo_dato": "2026-10-12 21:30:00",
    }
    base.update(k)
    return DatosPlanilla(**base)


@pytest.mark.unit
def test_nombre_con_estado_y_grupo_y_versiones(tmp_path):
    r1 = escribir_planilla(tmp_path, datos(), rec(), BLOQUES, False, 1)
    r2 = escribir_planilla(tmp_path, datos(), rec(), BLOQUES, False, 1)
    assert r1.name == "planilla_2026-10-13_OK_SOMBRA.xlsx"
    assert r2.name == "planilla_2026-10-13_OK_SOMBRA_v2.xlsx"
    assert nombre_disponible(tmp_path, "2026-10-13", "OK", "SOMBRA")[1] == 3
    assert not list(tmp_path.glob("*.tmp*"))


@pytest.mark.unit
def test_hojas_y_columnas(tmp_path):
    ruta = escribir_planilla(
        tmp_path, datos(advertencias=["Ruta 1: factor fuera de rango."]), rec(), BLOQUES, False, 1
    )
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Resumen", "Por bloque", "Por franja"]
    franja = pd.read_excel(ruta, sheet_name="Por franja")
    columnas = set(franja.columns)
    assert {
        "Ruta",
        "Bloque",
        "Franja",
        "Pasajeros predichos",
        "Meta de despachos",
        "Si hay conductor disponible",
    } <= columnas
    assert not any(c.startswith(("¿Siguió", "Motivo", "Despachos realizados")) for c in columnas)
    assert franja["Meta de despachos"].astype(str).tolist()[:3] == ["Sin meta", "3", "2"]
    bloque = pd.read_excel(ruta, sheet_name="Por bloque", nrows=3)  # sin las notas al pie
    assert bloque["Meta de despachos"].tolist() == [0, 3, 2]
    assert bloque["Pasajeros predichos"].tolist() == [12.3, 80.0, 60.0]
    resumen = pd.read_excel(ruta, sheet_name="Resumen").set_index("Campo")["Valor"]
    assert resumen["Estado"] == "OK" and resumen["Grupo"] == "SOMBRA"
    assert resumen["Advertencia"] == "Ruta 1: factor fuera de rango."
    assert "Corrida anticipada" not in resumen.index


@pytest.mark.unit
def test_corrida_anticipada_queda_marcada(tmp_path):
    ruta = escribir_planilla(tmp_path, datos(anticipada=True), rec(), BLOQUES, False, 1)
    resumen = pd.read_excel(ruta, sheet_name="Resumen").set_index("Campo")["Valor"]
    assert "incompletos" in resumen["Corrida anticipada"]


@pytest.mark.unit
def test_horario_de_bloques():
    assert horario_bloques(BLOQUES) == {
        "madrugada": "04:00 a 05:00",
        "pico_manana": "desde las 05:00",
    }
