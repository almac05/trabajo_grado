"""Orquestador: bloqueo, commit limpio, horario, estados, codigos de salida y corrida completa."""

import json
import time
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
import pytest
import yaml
from botocore.exceptions import ClientError, ReadTimeoutError
from openpyxl import load_workbook

from proyecto_grado.modeling.config import PROJECT_ROOT
from proyecto_grado.piloto.config import RUTA_CONFIG_PREDETERMINADA, sha256_archivo
from proyecto_grado.piloto.fechas import ZONA
from proyecto_grado.piloto.nocturno import (
    CODIGO_ERROR,
    CODIGO_OCUPADO,
    CODIGO_OK,
    Dependencias,
    Ocupado,
    bloqueo,
    cambios_que_ensucian,
    main,
)
from tests.piloto.conftest import requiere_semilla

LIMPIO = "?? analisis/nuevo.py\n?? comprobaciones.py\n"


def git_falso(porcelana=LIMPIO):
    def git(args):
        return "a" * 40 + "\n" if args[0] == "rev-parse" else porcelana

    return git


def reloj(h, m=0, dia=13):
    return lambda: datetime(2026, 10, dia, h, m, tzinfo=ZONA)


class FuenteQueFalla:
    nombre = "bd"

    def __init__(self, mensaje="se llego a la extraccion"):
        self.mensaje = mensaje

    def extraer(self, ventana, pk_ancla):
        raise RuntimeError(self.mensaje)


class EscritorFalso:
    def __init__(self, fallar_al_cerrar=False):
        self.log = []
        self.fallar_al_cerrar = fallar_al_cerrar

    def abrir_corrida(self, *a):
        self.log.append(("abrir", a))

    def cerrar_corrida(self, id_corrida, cierre):
        self.log.append(("cerrar", id_corrida, cierre))
        if self.fallar_al_cerrar:
            self.log.append(("error", (id_corrida,)))
            raise RuntimeError("la base no respondio al cerrar")

    def marcar_error(self, *a):
        self.log.append(("error", a))
        return True

    def cerrar(self):
        self.log.append(("fin",))


class CWFalso:
    """Cliente de CloudWatch: registra las llamadas o lanza el error indicado."""

    def __init__(self, error=None):
        self.llamadas = []
        self.error = error

    def put_metric_data(self, Namespace, MetricData):  # noqa: N803 - firma de boto3
        if self.error is not None:
            raise self.error
        self.llamadas.append({"Namespace": Namespace, "MetricData": MetricData})
        return {}

    def nombres(self):
        return [d["MetricName"] for ll in self.llamadas for d in ll["MetricData"]]


class S3Falso:
    def __init__(self):
        self.claves = []

    def put_object(self, Bucket, Key, Body):  # noqa: N803 - firma de boto3
        self.claves.append(Key)


@pytest.fixture
def entorno(monkeypatch):
    monkeypatch.setenv("PILOTO_ENTORNO", "LOCAL")


def config_minima(tmp_path: Path, **cambios) -> Path:
    """piloto.yaml con una sola entrada congelada: una semilla minima de dos despachos.

    tmp_path esta bajo la raiz del proyecto (--basetemp=.pytest_tmp), asi que la
    ruta relativa que exige la configuracion es valida.
    """
    semilla = tmp_path / "semilla.csv"
    semilla.write_text(
        "PK_INTERVALO_DESPACHO,FECHA_INICIAL\n100,2026-10-01\n200,2026-10-12\n", encoding="utf-8"
    )
    raw = yaml.safe_load(RUTA_CONFIG_PREDETERMINADA.read_text(encoding="utf-8"))
    raw["entradas"] = {
        "semilla_crudo": {
            "ruta": semilla.resolve().relative_to(PROJECT_ROOT.resolve()).as_posix(),
            "sha256": sha256_archivo(semilla),
        }
    }
    for clave, valor in cambios.items():
        seccion, campo = clave.split("__")
        raw[seccion][campo] = valor
    ruta = tmp_path / "piloto.yaml"
    ruta.write_text(yaml.safe_dump(raw, allow_unicode=True), encoding="utf-8")
    return ruta


def correr(tmp_path, argv, deps, config=None):
    salidas = tmp_path / "salidas"
    config = config or config_minima(tmp_path)
    codigo = main([*argv, "--directorio-salidas", str(salidas), "--config", str(config)], deps)
    manifiestos = sorted(salidas.glob("ejecuciones/*/manifiesto.json"))
    m = json.loads(manifiestos[-1].read_text(encoding="utf-8")) if manifiestos else None
    return codigo, m, salidas


def deps(**k):
    base = {
        "reloj": reloj(0, 30),
        "git": git_falso(),
        "cargar_env": lambda: None,
        "fuente": lambda cfg, nombre: FuenteQueFalla(),
    }
    base.update(k)
    return Dependencias(**base)


# ------------------------------------------------------------------ git


@pytest.mark.unit
def test_regla_de_commit_limpio():
    porcelana = (
        " M src/proyecto_grado/x.py\nM  README.md\n?? src/nuevo.py\n?? configs/piloto/x.yaml\n"
        "?? scripts/piloto/a.ps1\n?? analisis/41.py\n?? comprobaciones.py\n?? docs/nota.md\n"
    )
    assert cambios_que_ensucian(porcelana) == [
        " M src/proyecto_grado/x.py",
        "M  README.md",
        "?? src/nuevo.py",
        "?? configs/piloto/x.yaml",
        "?? scripts/piloto/a.ps1",
    ]
    assert cambios_que_ensucian(LIMPIO) == []


@pytest.mark.unit
def test_commit_sucio_no_corre_y_queda_registrado(tmp_path, entorno):
    codigo, m, salidas = correr(tmp_path, [], deps(git=git_falso(" M src/a.py\n")))
    assert codigo == CODIGO_ERROR
    assert "commit limpio" in m["error"]["mensaje"]
    assert m["codigo"]["status_porcelain"] == " M src/a.py\n"
    assert "extraccion" not in m["pasos"]
    assert "Estado: ERROR" in (salidas / "ESTADO_ULTIMA_EJECUCION.txt").read_text(encoding="utf-8")


@pytest.mark.unit
def test_permitir_sucio_corre_y_queda_registrado(tmp_path, entorno):
    _, m, _ = correr(tmp_path, ["--permitir-sucio"], deps(git=git_falso(" M src/a.py\n")))
    assert m["codigo"]["permitir_sucio"] is True and "extraccion" in m["pasos"]
    assert any("permitir-sucio" in a for a in m["advertencias"])


# ------------------------------------------------------------------ bloqueo


@pytest.mark.unit
def test_bloqueo_exclusivo_y_vencido(tmp_path):
    ruta = tmp_path / "x.lock"
    ahora = datetime(2026, 10, 13, 0, 30, tzinfo=ZONA)
    with bloqueo(ruta, ahora):
        with pytest.raises(Ocupado), bloqueo(ruta, ahora):
            pass
        assert ruta.exists()
    assert not ruta.exists()
    ruta.write_text(json.dumps({"pid": 1, "inicio": (ahora - timedelta(hours=7)).isoformat()}))
    with bloqueo(ruta, ahora):  # abandonado: se reemplaza
        pass


@pytest.mark.unit
def test_segunda_ejecucion_termina_con_codigo_2(tmp_path, entorno):
    salidas = tmp_path / "salidas"
    salidas.mkdir()
    (salidas / "nocturno.lock").write_text(
        json.dumps({"pid": 1, "inicio": datetime(2026, 10, 13, 0, 29, tzinfo=ZONA).isoformat()})
    )
    codigo, m, _ = correr(tmp_path, [], deps())
    assert codigo == CODIGO_OCUPADO and m is None
    assert "ya hay otra en curso" in (salidas / "ejecuciones_rechazadas.log").read_text("utf-8")
    assert (salidas / "nocturno.lock").exists()  # el de la otra ejecucion no se toca


# ------------------------------------------------------------------ horario y entorno


@pytest.mark.unit
@pytest.mark.parametrize(
    ("hora", "permitido"), [((3, 59), True), ((4, 0), False), ((20, 59), False), ((21, 0), True)]
)
def test_horario_sin_base(tmp_path, entorno, hora, permitido):
    codigo, m, _ = correr(tmp_path, [], deps(reloj=reloj(*hora)))
    assert codigo == CODIGO_ERROR  # la fuente simulada falla al extraer
    assert m["horario"]["horario_permitido"] is permitido
    if permitido:
        assert "se llego a la extraccion" in m["error"]["mensaje"]
    else:
        assert "no se consulta la base" in m["error"]["mensaje"]
        assert "extraccion" not in m["pasos"]


@pytest.mark.unit
def test_forzar_horario_queda_registrado(tmp_path, entorno):
    _, m, _ = correr(tmp_path, ["--forzar-horario"], deps(reloj=reloj(10)))
    assert m["horario"]["forzar_horario"] is True and "extraccion" in m["pasos"]
    assert any("forzar-horario" in a for a in m["advertencias"])


@pytest.mark.unit
def test_fuente_local_no_esta_sujeta_al_horario(tmp_path, entorno):
    _, m, _ = correr(tmp_path, ["--fuente", "local"], deps(reloj=reloj(10)))
    assert m["horario"]["usa_bd"] is False and "extraccion" in m["pasos"]


@pytest.mark.unit
def test_corrida_anticipada_y_tardia_quedan_marcadas(tmp_path, entorno):
    _, m, _ = correr(tmp_path, [], deps(reloj=reloj(22)))
    assert m["objetivo"]["anticipada"] and m["objetivo"]["d_mas_1"] == "2026-10-14"
    assert any("anticipada" in a for a in m["advertencias"])
    _, m, _ = correr(tmp_path, [], deps(reloj=reloj(3, 59, dia=14)))
    assert not m["objetivo"]["tardia"]
    _, m, _ = correr(tmp_path, ["--fuente", "local"], deps(reloj=reloj(6, dia=14)))
    assert m["objetivo"]["tardia"] and m["objetivo"]["d_mas_1"] == "2026-10-14"


@pytest.mark.unit
def test_falta_piloto_entorno(tmp_path, monkeypatch):
    monkeypatch.delenv("PILOTO_ENTORNO", raising=False)
    codigo, m, _ = correr(tmp_path, [], deps())
    assert codigo == CODIGO_ERROR and "PILOTO_ENTORNO" in m["error"]["mensaje"]


@pytest.mark.unit
def test_entrada_congelada_alterada_es_error(tmp_path, entorno):
    raw = yaml.safe_load(RUTA_CONFIG_PREDETERMINADA.read_text(encoding="utf-8"))
    # Las rutas de entrada son relativas a la raiz del proyecto: una que no existe.
    raw["entradas"] = {"x": {"ruta": "no/existe/x.csv", "sha256": "0" * 64}}
    config = tmp_path / "piloto.yaml"
    config.write_text(yaml.safe_dump(raw), encoding="utf-8")
    codigo, m, _ = correr(tmp_path, [], deps(), config=config)
    assert codigo == CODIGO_ERROR and "ErrorEntradaCongelada" in m["error"]["mensaje"]


@pytest.mark.unit
def test_secretos_ocultos_en_manifiesto_log_y_estado(tmp_path, entorno, monkeypatch):
    monkeypatch.setenv("REGISTEL_DB_PASSWORD", "clave-ultra-secreta")
    d = deps(fuente=lambda cfg, n: FuenteQueFalla("fallo con clave-ultra-secreta"))
    codigo, m, salidas = correr(tmp_path, [], d)
    assert codigo == CODIGO_ERROR and "[OCULTO]" in m["error"]["mensaje"]
    carpeta = next(salidas.glob("ejecuciones/*"))
    for archivo in [
        carpeta / "manifiesto.json",
        carpeta / "ejecucion.log",
        salidas / "ESTADO_ULTIMA_EJECUCION.txt",
    ]:
        texto = archivo.read_text(encoding="utf-8")
        assert "clave-ultra-secreta" not in texto and "[OCULTO]" in texto


@pytest.mark.unit
def test_error_con_corrida_abierta_la_marca_en_error(tmp_path, entorno):
    escritor = EscritorFalso()
    codigo, m, _ = correr(tmp_path, ["--escribir-bd"], deps(escritor_bd=lambda: escritor))
    assert codigo == CODIGO_ERROR
    assert [x[0] for x in escritor.log] == ["abrir", "error", "fin"]
    assert m["bd"]["marcada_error"] is True


@pytest.mark.unit
def test_argumentos_excluyentes():
    with pytest.raises(SystemExit):
        main(["--relleno", "2026-10-04", "2026-10-05", "--fecha-objetivo", "2026-10-05"])
    with pytest.raises(SystemExit):
        main(["--relleno", "2026-10-05", "2026-10-04"])
    with pytest.raises(SystemExit):
        main(["--relleno", "2026-10-04", "2026-10-05", "--escribir-bd"])
    for otra in (["--relleno", "2026-10-04", "2026-10-05"], ["--sincronizar-estado"]):
        with pytest.raises(SystemExit):
            main([*otra, "--publicar-metricas"])


# ------------------------------------------------------------------ metricas de CloudWatch


@pytest.fixture
def metricas_activadas(monkeypatch):
    monkeypatch.setenv("PILOTO_METRICAS", "1")


def fabrica(cw, creados=None):
    def crear():
        if creados is not None:
            creados.append(1)
        return cw

    return crear


@pytest.mark.unit
def test_ensayo_sin_publicar_metricas_no_publica(tmp_path, entorno, metricas_activadas):
    creados = []
    argv = ["--fecha-objetivo", "2026-10-01", "--fuente", "local"]
    codigo, m, _ = correr(tmp_path, argv, deps(cliente_cw=fabrica(CWFalso(), creados)))
    assert codigo == CODIGO_ERROR and creados == []
    assert m["metricas"] == {
        "metricas_publicadas": False,
        "motivo": "no_aplica: ensayo sin --publicar-metricas",
    }


@pytest.mark.unit
def test_ensayo_con_publicar_metricas_publica(tmp_path, entorno, metricas_activadas):
    cw = CWFalso()
    argv = ["--fecha-objetivo", "2026-10-01", "--fuente", "local", "--publicar-metricas"]
    codigo, m, _ = correr(tmp_path, argv, deps(cliente_cw=fabrica(cw)))
    assert codigo == CODIGO_ERROR  # la fuente simulada falla al extraer
    assert m["metricas"]["metricas_publicadas"] is True
    assert cw.llamadas[0]["Namespace"] == "Montebello/Despacho"
    datos = {d["MetricName"]: d for d in cw.llamadas[0]["MetricData"]}
    assert datos["EstadoCorrida"]["Value"] == 3  # ERROR
    assert datos["DuracionTotal"]["Unit"] == "Seconds"
    assert datos["DuracionExtraccion"]["Value"] >= 0  # el paso llego a empezar
    assert "RegistrosExtraidos" not in datos and "FactorCorreccion" not in datos
    assert datos["EstadoCorrida"]["Timestamp"].utcoffset().total_seconds() == 0
    assert m["metricas"]["metricas"] == [d["MetricName"] for d in cw.llamadas[0]["MetricData"]]


@pytest.mark.unit
def test_corrida_nocturna_publica_tambien_en_error(tmp_path, entorno, metricas_activadas):
    cw = CWFalso()
    d = deps(cliente_cw=fabrica(cw), git=git_falso(" M src/a.py\n"))
    codigo, m, _ = correr(tmp_path, [], d)
    assert codigo == CODIGO_ERROR and m["modo"] == "nocturno"
    assert cw.nombres()[:2] == ["EstadoCorrida", "DuracionTotal"]
    assert "DuracionExtraccion" not in cw.nombres()  # error antes de la extraccion
    assert m["metricas"]["metricas_publicadas"] is True


@pytest.mark.unit
@pytest.mark.parametrize(
    "error",
    [
        ClientError({"Error": {"Code": "AccessDenied", "Message": "x"}}, "PutMetricData"),
        ReadTimeoutError(endpoint_url="https://monitoring.us-east-2.amazonaws.com"),
    ],
    ids=["client_error", "timeout"],
)
def test_fallo_de_cloudwatch_no_cambia_estado_ni_codigo(
    tmp_path, entorno, metricas_activadas, error
):
    codigo, m, salidas = correr(tmp_path, [], deps(cliente_cw=fabrica(CWFalso(error))))
    assert codigo == CODIGO_ERROR and m["estado"] == "ERROR"
    assert "se llego a la extraccion" in m["error"]["mensaje"]  # el error es el de la corrida
    assert m["metricas"]["metricas_publicadas"] is False
    assert type(error).__name__ in m["metricas"]["motivo"]
    assert "Estado: ERROR" in (salidas / "ESTADO_ULTIMA_EJECUCION.txt").read_text("utf-8")
    log = next(salidas.glob("ejecuciones/*/ejecucion.log")).read_text(encoding="utf-8")
    assert "No se pudieron publicar las metricas" in log


@pytest.mark.unit
def test_sin_piloto_metricas_no_se_crea_cliente(tmp_path, entorno, monkeypatch):
    monkeypatch.delenv("PILOTO_METRICAS", raising=False)
    creados = []
    _, m, _ = correr(tmp_path, [], deps(cliente_cw=fabrica(CWFalso(), creados)))
    assert creados == [] and m["metricas"]["motivo"].startswith("desactivadas")


@pytest.mark.unit
def test_codigo_2_no_publica(tmp_path, entorno, metricas_activadas):
    salidas = tmp_path / "salidas"
    salidas.mkdir()
    (salidas / "nocturno.lock").write_text(
        json.dumps({"pid": 1, "inicio": datetime(2026, 10, 13, 0, 29, tzinfo=ZONA).isoformat()})
    )
    creados = []
    codigo, _, _ = correr(tmp_path, [], deps(cliente_cw=fabrica(CWFalso(), creados)))
    assert codigo == CODIGO_OCUPADO and creados == []


# ------------------------------------------------------------------ integracion


@requiere_semilla
@pytest.mark.integration
@pytest.mark.slow
def test_corrida_completa_sin_base_de_datos(tmp_path, entorno):
    """Orquestador de punta a punta sobre el 1 de octubre con la semilla como extraccion."""
    escritor, s3 = EscritorFalso(), S3Falso()
    cw = CWFalso(ClientError({"Error": {"Code": "AccessDenied", "Message": "x"}}, "PutMetricData"))
    d = Dependencias(
        reloj=reloj(0, 30, dia=13),
        git=git_falso(),
        cargar_env=lambda: None,
        escritor_bd=lambda: escritor,
        cliente_s3=lambda: s3,
        cliente_cw=lambda: cw,
    )
    salidas = tmp_path / "salidas"
    t0 = time.perf_counter()
    import os

    os.environ["PILOTO_S3_BUCKET"] = "bucket-prueba"
    os.environ["PILOTO_METRICAS"] = "1"
    try:
        codigo = main(
            [
                "--fecha-objetivo",
                "2026-10-01",
                "--fuente",
                "local",
                "--escribir-bd",
                "--subir-s3",
                "--publicar-metricas",
                "--directorio-salidas",
                str(salidas),
            ],
            d,
        )
    finally:
        os.environ.pop("PILOTO_S3_BUCKET", None)
        os.environ.pop("PILOTO_METRICAS", None)
    segundos = time.perf_counter() - t0
    print(f"\nCorrida de integracion: {segundos:.1f} s")
    assert codigo == CODIGO_OK

    planilla = salidas / "planillas" / "planilla_2026-10-01_OK_SOMBRA.xlsx"
    assert load_workbook(planilla).sheetnames == ["Resumen", "Por bloque", "Por franja"]
    carpeta = next(salidas.glob("ejecuciones/*"))
    m = json.loads((carpeta / "manifiesto.json").read_text(encoding="utf-8"))
    for clave in (
        "codigo",
        "entradas",
        "extraccion",
        "serie",
        "prediccion",
        "factor",
        "recomendacion",
        "salidas",
        "pasos",
        "s3",
        "bd",
    ):
        assert clave in m, clave
    assert m["estado"] == "OK" and m["modo"] == "ensayo" and m["grupo"] == "SOMBRA"
    # CloudWatch caido no cambia el estado ni el codigo; el desempeno de D (30 sep) no
    # tiene prediccion emitida por la piloto (el historial de un ensayo es el backtest).
    assert m["metricas"]["metricas_publicadas"] is False
    assert "AccessDenied" in m["metricas"]["motivo"]
    assert m["desempeno"]["jornada"] == "2026-09-30"
    assert "sin_prediccion_emitida" in m["desempeno"]["motivo"]
    assert m["recomendacion"]["filas"] == 57
    assert m["extraccion"]["pk_ancla"] > 0 and m["serie"]["sha256"]
    assert set(m["factor"]) == {"1", "3"}
    assert all(len(v["sha256"]) == 64 for v in m["entradas"].values())
    # Ensayo con fuente local: no toca el puntero del almacen ni el historial.
    assert not (salidas / "almacen" / "ACTUAL.json").exists()
    assert not (salidas / "historial").exists()

    acciones = [x[0] for x in escritor.log]
    assert acciones == ["abrir", "cerrar", "fin"]
    cierre = escritor.log[1][2]
    assert cierre.estado == "OK" and len(cierre.filas) == 57
    assert cierre.ruta_s3 == f"s3://bucket-prueba/ejecuciones/{m['id_ejecucion']}/"
    assert cierre.version_dataset == m["serie"]["sha256"]
    # La planilla sube solo al final, despues del cierre en la base, con el manifiesto final.
    assert s3.claves[-3:] == [
        "planillas/planilla_2026-10-01_OK_SOMBRA.xlsx",
        f"ejecuciones/{m['id_ejecucion']}/manifiesto.json",
        f"ejecuciones/{m['id_ejecucion']}/ejecucion.log",
    ]
    assert sum(k.startswith("planillas/") for k in s3.claves) == 1
    pasos = list(m["pasos"])
    assert pasos.index("s3") < pasos.index("bd_cerrar") < pasos.index("planilla")

    pred = pd.read_csv(carpeta / "predicciones.csv")
    x43 = pd.read_csv(
        "analisis/salidas/predicciones__xgboost__correccion_nivel_20251127_20261003.csv"
    )
    u = pred.merge(x43[x43["dia_op"] == "2026-10-01"], on=["route", "hora"], suffixes=("", "_43"))
    assert len(u) == 51 and (u["pred_xgboost"] - u["pred_xgboost_43"]).abs().max() < 1e-4


@requiere_semilla
@pytest.mark.integration
@pytest.mark.slow
def test_fallo_al_cerrar_en_la_base_emite_solo_planilla_error(tmp_path, entorno, monkeypatch):
    """Ninguna planilla (local ni en S3) puede tener un estado distinto al de la corrida."""
    escritor, s3, cw = EscritorFalso(fallar_al_cerrar=True), S3Falso(), CWFalso()
    d = Dependencias(
        reloj=reloj(0, 30, dia=13),
        git=git_falso(),
        cargar_env=lambda: None,
        escritor_bd=lambda: escritor,
        cliente_s3=lambda: s3,
        cliente_cw=lambda: cw,
    )
    monkeypatch.setenv("PILOTO_S3_BUCKET", "bucket-prueba")
    monkeypatch.setenv("PILOTO_METRICAS", "1")
    salidas = tmp_path / "salidas"
    codigo = main(
        [
            "--fecha-objetivo",
            "2026-10-01",
            "--fuente",
            "local",
            "--escribir-bd",
            "--subir-s3",
            "--publicar-metricas",
            "--directorio-salidas",
            str(salidas),
        ],
        d,
    )
    assert codigo == CODIGO_ERROR
    datos = {
        (x["MetricName"], tuple(v["Value"] for v in x.get("Dimensions", []))): x["Value"]
        for x in cw.llamadas[0]["MetricData"]
    }
    assert datos[("EstadoCorrida", ())] == 3
    assert datos[("FilasRecomendaciones", ())] == 57  # generadas (la base no cerro)
    assert {("FactorCorreccion", ("1",)), ("FactorCorreccion", ("3",))} <= set(datos)
    assert datos[("RegistrosExtraidos", ())] > 0
    planillas = sorted(p.name for p in (salidas / "planillas").glob("*.xlsx"))
    assert planillas == ["planilla_2026-10-01_ERROR_SOMBRA.xlsx"]
    resumen = pd.read_excel(salidas / "planillas" / planillas[0], sheet_name="Resumen")
    resumen = resumen.set_index("Campo")["Valor"]
    assert resumen["Estado"] == "ERROR"
    assert any("no respondio" in str(v) for v in resumen[resumen.index == "Advertencia"])
    assert [k for k in s3.claves if k.startswith("planillas/")] == [
        "planillas/planilla_2026-10-01_ERROR_SOMBRA.xlsx"
    ]
    m = json.loads(next(salidas.glob("ejecuciones/*/manifiesto.json")).read_text("utf-8"))
    assert m["estado"] == "ERROR" and m["salidas"]["planilla"]["ruta"].endswith(
        "_ERROR_SOMBRA.xlsx"
    )
    assert "historial" not in m and not (salidas / "historial").exists()


@pytest.mark.unit
def test_consolidar_estado_solo_con_ok(tmp_path):
    from proyecto_grado.piloto.nocturno import Datos, Ejecucion, Estado, consolidar_estado

    alm = tmp_path / "almacen"
    alm.mkdir()
    (alm / "crudo_x.parquet").write_bytes(b"c")
    (alm / "conjunto_apto_x.parquet").write_bytes(b"a")
    datos = Datos(
        conjunto_apto=pd.DataFrame(),
        instantanea=tmp_path / "i.parquet",
        almacen_dir=alm,
        almacen_rutas={
            "crudo": alm / "crudo_x.parquet",
            "conjunto_apto": alm / "conjunto_apto_x.parquet",
        },
        almacen_cubierto=date(2026, 10, 12),
        almacen_oficial=True,
    )
    filas = pd.DataFrame(
        {
            "id_ejecucion": ["x"],
            "origen": ["piloto"],
            "version": ["definitiva"],
            "emitido": [""],
            "route": [1],
            "dia_op": ["2026-10-13"],
            "hora": ["05:00"],
            "tipo_dia": ["LABORAL"],
            "regimen": ["ordinario"],
            "pred_xgboost": [1.0],
            "ref52": [1.0],
            "factor": [1.0],
            "pred_final": [1.0],
        }
    )
    e = Ejecucion(
        id="x",
        inicio=datetime(2026, 10, 13, tzinfo=ZONA),
        carpeta=tmp_path,
        salidas=tmp_path,
        manifiesto={"almacen": {}},
    )
    e.estado = Estado.ERROR
    consolidar_estado(e, datos, filas)
    assert not (alm / "ACTUAL.json").exists() and not (tmp_path / "historial").exists()
    e.estado = Estado.OK
    consolidar_estado(e, datos, filas)
    assert json.loads((alm / "ACTUAL.json").read_text())["cubierto_hasta"] == "2026-10-12"
    assert (tmp_path / "historial" / "ACTUAL.json").exists()
    assert e.manifiesto["almacen"]["puntero_actualizado"] is True


@requiere_semilla
@pytest.mark.integration
@pytest.mark.slow
def test_relleno_encadena_el_factor(tmp_path, entorno):
    """Relleno 1-2 oct con backtest hasta el 30 sep: el factor del 2 usa el relleno del 1."""
    config = tmp_path / "piloto.yaml"
    raw = yaml.safe_load(RUTA_CONFIG_PREDETERMINADA.read_text(encoding="utf-8"))
    raw["correccion"]["fin_backtest"] = date(2026, 9, 30)
    config.write_text(yaml.safe_dump(raw), encoding="utf-8")
    d = Dependencias(reloj=reloj(0, 30, dia=13), git=git_falso(), cargar_env=lambda: None)
    codigo, m, salidas = correr(
        tmp_path, ["--relleno", "2026-10-01", "2026-10-02", "--fuente", "local"], d, config=config
    )
    assert codigo == CODIGO_OK, m.get("error")
    assert m["modo"] == "relleno" and m["relleno"]["escrito_en_historial"] is False
    rel = pd.read_csv(next(salidas.glob("ejecuciones/*")) / "relleno.csv")
    assert set(rel["origen"]) == {"relleno"} and set(rel["dia_op"]) == {"2026-10-01", "2026-10-02"}
    s44 = pd.read_csv("analisis/salidas/serie_correccion__correccion_nivel_20251127_20261003.csv")
    s44 = s44[(s44["ventana_dias"] == 14)].set_index(["dia_op", "route"])["factor"]
    for dia in ("2026-10-01", "2026-10-02"):
        for ruta in ("1", "3"):
            factor = m["relleno"]["jornadas"][dia]["factores"][ruta]
            assert factor == pytest.approx(s44.loc[(dia, int(ruta))], abs=1e-5)
